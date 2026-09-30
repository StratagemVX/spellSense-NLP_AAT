import docx
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.8)
sec.bottom_margin = Inches(0.8)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

header = sec.header
header.is_linked_to_previous = False
p_head = header.paragraphs[0]
run_head = p_head.add_run()
run_head.add_picture("report_assets/page1_img0_1486x1985.png", width=Inches(7.8), height=Inches(10.8))

# Find the inline drawing element and transform to floating behindDoc
inline_elems = run_head._r.xpath('.//wp:inline')
if inline_elems:
    inline = inline_elems[0]
    extent = inline.find(docx.oxml.ns.qn('wp:extent'))
    docPr = inline.find(docx.oxml.ns.qn('wp:docPr'))
    graphic = inline.find(docx.oxml.ns.qn('a:graphic'))
    cx = extent.get('cx')
    cy = extent.get('cy')

    anchor_xml = f"""
    <wp:anchor {nsdecls('wp')} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
      <wp:simplePos x="0" y="0"/>
      <wp:positionH relativeFrom="page">
        <wp:align>center</wp:align>
      </wp:positionH>
      <wp:positionV relativeFrom="page">
        <wp:align>center</wp:align>
      </wp:positionV>
      <wp:extent cx="{cx}" cy="{cy}"/>
      <wp:effectExtent l="0" t="0" r="0" b="0"/>
      <wp:wrapNone/>
      <wp:docPr id="101" name="CoverBorder"/>
      <wp:cNvGraphicFramePr/>
    </wp:anchor>
    """
    anchor = parse_xml(anchor_xml)
    anchor.append(graphic)
    drawing = inline.getparent()
    drawing.replace(inline, anchor)

p = doc.add_paragraph("Testing text inside border")
doc.save("test_border.docx")
print("Saved test_border.docx")
