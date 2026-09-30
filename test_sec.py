from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = Document()

def set_section_page_numbering(section, fmt="decimal", start=1):
    sectPr = section._sectPr
    pgNumTypes = sectPr.xpath('./w:pgNumType')
    if pgNumTypes:
        pgNumType = pgNumTypes[0]
    else:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    if fmt:
        pgNumType.set(qn('w:fmt'), fmt)
    if start is not None:
        pgNumType.set(qn('w:start'), str(start))

def add_page_number_field(run):
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    run._r.append(fldSimple)

# Sec 1
s1 = doc.sections[0]
s1.header.is_linked_to_previous = False
s1.footer.is_linked_to_previous = False
p1 = doc.add_paragraph("Cover Page")

# Sec 2: Prelim
s2 = doc.add_section()
s2.header.is_linked_to_previous = False
s2.footer.is_linked_to_previous = False
set_section_page_numbering(s2, fmt="lowerRoman", start=1)
r2 = s2.footer.paragraphs[0].add_run("Prelim Page: ")
add_page_number_field(r2)
doc.add_paragraph("Acknowledgement")
doc.add_page_break()
doc.add_paragraph("Abstract")

# Sec 3: Main
s3 = doc.add_section()
s3.header.is_linked_to_previous = False
s3.footer.is_linked_to_previous = False
set_section_page_numbering(s3, fmt="decimal", start=1)
r3 = s3.footer.paragraphs[0].add_run("Main Page: ")
add_page_number_field(r3)
doc.add_paragraph("Chapter 1")
doc.add_page_break()
doc.add_paragraph("Chapter 2")

doc.save("test_sections.docx")
print("Saved test_sections.docx")
